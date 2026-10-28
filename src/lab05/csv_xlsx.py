import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    '''Преобразует CSV в XLSX.
    Первая строка CSV — заголовок. Лист называется "Sheet1".
    Ширина колонок — авто-ширина по тексту (не менее 8).
    '''
    source = Path(csv_path)
    target = Path(xlsx_path)

    # 1. Валидация файлов
    if not source.is_file():
        raise FileNotFoundError(f"Файл не найден: {csv_path}")
    if source.suffix.lower() != ".csv":
        raise ValueError("Входной файл должен быть CSV")
    if target.suffix.lower() != ".xlsx":
        raise ValueError("Выходной файл должен быть XLSX")
    
    # 2. Чтение CSV
    rows = []
    try:
        with source.open(encoding="utf-8", newline="") as f:
            for row in csv.reader(f):
                if row: # Пропускаем пустые строки
                    rows.append(row)
    except Exception as e:
        raise ValueError(f"Ошибка чтения CSV: {e}")

    # 3. Проверка на наличие данных (заголовок + хотя бы одна строка)
    if len(rows) < 2:
        raise ValueError("В CSV нет данных (только заголовок или файл пуст)")

    # 4. Запись в Excel
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        
        for row in rows:
            ws.append(row)

        # 5. Авто-ширина колонок
        widths = {}
        for row in rows:
            for i, value in enumerate(row, start=1):
                # str() на случай, если попадутся не строки
                widths[i] = max(widths.get(i, 0), len(str(value)))

        for i, width in widths.items():
            ws.column_dimensions[get_column_letter(i)].width = max(width, 8)

        # Создаем папку, если её нет, и сохраняем
        target.parent.mkdir(parents=True, exist_ok=True)
        wb.save(xlsx_path)
        
    except Exception as e:
        raise ValueError(f"Ошибка при создании XLSX файла: {e}")